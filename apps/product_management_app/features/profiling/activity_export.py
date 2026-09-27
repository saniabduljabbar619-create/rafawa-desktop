# apps/product_management_app/features/profiling/activity_export.py
# -*- coding: utf-8 -*-
"""
Exports the Customer Profile Activity tab's rows to a real .xlsx file -
everything currently shown on screen: every transaction (sale, purchase,
advance, return), every individual payment, and every settlement marker,
in the same order and with the same column layout already visible in
the app, so what the user sees and what they export always match.
"""
from __future__ import annotations

from typing import List, Dict, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = ["Date", "Ledger", "Type", "Reference", "Summary",
          "Total", "Paid", "Balance", "Balance Before", "Balance After",
          "Status", "Operator"]


def export_activity_to_excel(acts: List[Dict], party_name: str, output_path: str,
                             operator_lookup=None) -> None:
    """
    acts: the exact list PartyLedgerService.list_activity_details() (or
    list_walkin_activity()) returns - the same data load_activity_table()
    in customer_profile.py already renders, so the export always matches
    what's on screen.

    operator_lookup: optional callable(reference_str) -> str, used the
    same way customer_profile.py's own _fetch_operator_by_reference is,
    for rows that don't already carry their own "operator" key directly
    (payment rows do; sale/purchase/advance/return rows don't).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Activity"

    title_font = Font(bold=True, size=13)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2E7D46")
    marker_font = Font(bold=True, size=12, color="1E8C46")
    money_font = Font(bold=True)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    ws["A1"] = f"Activity - {party_name}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center")

    header_row = 3
    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=c, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    row_idx = header_row + 1
    for a in acts:
        if a.get("type") == "settlement_marker":
            text = str(a.get("summary", ""))
            display = f"---- {text} ----"

            ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=len(COLUMNS))
            cell = ws.cell(row=row_idx, column=1, value=display)
            cell.font = marker_font
            cell.alignment = Alignment(horizontal="center")
            row_idx += 1
            continue

        ref_str = a.get("reference", "")
        if a.get("type") == "payment" and a.get("operator"):
            operator = a["operator"]
        elif operator_lookup:
            operator = operator_lookup(ref_str)
        else:
            operator = ""

        total_val = a.get("total")
        if total_val is None:
            total_val = a.get("amount")
        total_val = float(total_val or 0)

        paid_val = a.get("paid")
        balance_val = a.get("balance")
        bal_before = a.get("balance_before")
        bal_after = a.get("balance_after")
        status = str(a.get("status", "") or "")

        values = [
            str(a.get("dt") or "")[:19],
            a.get("ledger", ""),
            a.get("type", ""),
            ref_str,
            a.get("summary", ""),
            total_val,
            float(paid_val) if paid_val is not None else "",
            float(balance_val) if balance_val is not None else "",
            float(bal_before) if bal_before is not None else "",
            float(bal_after) if bal_after is not None else "",
            status.capitalize(),
            operator,
        ]
        for c, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=c, value=val)
            if c == 6:  # Total
                cell.font = money_font
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0.00"
            elif c in (7, 8, 9, 10) and isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"
        row_idx += 1

    # Reasonable column widths, matching roughly what the on-screen
    # table's own proportions look like.
    widths = [20, 12, 12, 20, 22, 14, 14, 14, 14, 14, 12, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = f"A{header_row + 1}"

    wb.save(output_path)