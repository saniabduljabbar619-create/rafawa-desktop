from datetime import date, timedelta

import openpyxl

from database.models.sale import Sale
from database.services.report_export import export_reports
from database.services.report_service import ReportService
from database.services.return_service import ReturnService
from database.services.sales_service import SaleService
from database.session import get_session
from tests.conftest import ADMIN_ID, CASHIER_A, CASHIER_B, CROP_ID, CUSTOMER_ID, FERT_ID, make_sale, pay


def test_period_summary_adds_up():
    a = make_sale(user=CASHIER_A)                                  # 20000
    make_sale(user=CASHIER_B)                                      # 20000
    pay("sale", a["sale_id"], CUSTOMER_ID, "in", 5000, user=CASHIER_A, method="cash")
    pay("sale", a["sale_id"], CUSTOMER_ID, "in", 3000, user=CASHIER_B, method="transfer")
    ReturnService.create_return(document_type="sale", document_id=a["sale_id"],
                                item_id=a["sale_item_ids"][0], qty_kg=100, qty_bags=1,
                                reason="x", created_by_user_id=CASHIER_A)          # -10000
    rep = ReportService.period_summary()
    assert rep["sales"]["count"] == 2 and rep["sales"]["gross"] == 40000
    assert rep["sales"]["returns"] == 10000 and rep["sales"]["net"] == 30000
    assert rep["cash"] == {"in": 8000.0, "out": 0.0, "net": 8000.0}
    assert {m["method"]: m["in"] for m in rep["by_method"]} == {"cash": 5000.0, "transfer": 3000.0}
    cashiers = {c["username"]: c for c in rep["by_cashier"]}
    assert cashiers["alice"]["sales"] == 20000 and cashiers["alice"]["cash_in"] == 5000
    assert cashiers["bob"]["cash_in"] == 3000


def test_period_summary_respects_dates_and_ignores_voided():
    keep = make_sale()
    gone = make_sale()
    SaleService.void_sale(gone["sale_id"], ADMIN_ID)
    assert ReportService.period_summary()["sales"]["count"] == 1
    today = date.today()
    assert ReportService.period_summary(today, today)["sales"]["count"] == 1
    assert ReportService.period_summary(today + timedelta(days=1), None)["sales"]["count"] == 0
    with get_session() as s:                                         # move it 40 days back
        from datetime import datetime
        s.get(Sale, keep["sale_id"]).created_at = datetime.now() - timedelta(days=40)
    assert ReportService.period_summary(today - timedelta(days=7), today)["sales"]["count"] == 0


def test_receivables_aging_buckets_and_exclusions():
    from datetime import datetime
    old = make_sale()                                                # 20000, will be 70 days old
    new = make_sale()                                                # 20000 today
    paid = make_sale()
    pay("sale", paid["sale_id"], CUSTOMER_ID, "in", 20000)           # fully paid: excluded
    voided = make_sale()
    SaleService.void_sale(voided["sale_id"], ADMIN_ID)               # voided: excluded
    pay("sale", old["sale_id"], CUSTOMER_ID, "in", 5000)
    with get_session() as s:
        s.get(Sale, old["sale_id"]).created_at = datetime.now() - timedelta(days=70)
    rows = ReportService.aging("sale")
    assert len(rows) == 1
    r = rows[0]
    assert r["party_id"] == CUSTOMER_ID and r["documents"] == 2
    assert r["0-30"] == 20000 and r["61-90"] == 15000 and r["31-60"] == 0 and r["90+"] == 0
    assert r["total"] == 35000


def test_stock_valuation_values_at_cost():
    make_sale()                                                      # leaves 18 bags / 1800 kg maize
    rows = {r["product"]: r for r in ReportService.stock_valuation()}
    assert rows["Maize"]["kg"] == 1800 and rows["Maize"]["unit"] == "kg" and rows["Maize"]["value"] == 1800 * 80
    assert rows["NPK"]["bags"] == 50 and rows["NPK"]["value"] == 50 * 9000


def test_excel_export_opens_and_has_the_numbers(tmp_path):
    a = make_sale()
    pay("sale", a["sale_id"], CUSTOMER_ID, "in", 5000)
    out = export_reports(str(tmp_path / "r.xlsx"))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Summary", "By cashier", "By payment method", "Receivables", "Payables",
                             "Stock valuation"]
    summary = {r[0].value: r[1].value for r in wb["Summary"].iter_rows(min_row=5) if r[0].value}
    assert summary["Sales - gross"] == 20000 and summary["Cash received"] == 5000
    recv = [[c.value for c in r] for r in wb["Receivables"].iter_rows(min_row=5)]
    assert recv[0][0] == "Customer One" and recv[0][-1] == 15000
    assert recv[-1][0] == "TOTAL" and recv[-1][-1] == 15000
    assert wb["Stock valuation"].freeze_panes == "A5"


def test_aging_rejects_bad_kind():
    import pytest
    with pytest.raises(ValueError):
        ReportService.aging("advance")


def test_reports_are_admin_only():
    import pytest
    from shared.auth.permissions import PermissionDenied
    from shared.auth.session import UserSession, set_session
    set_session(UserSession(user_id=CASHIER_A, username="alice", role="cashier"))
    for call in (ReportService.period_summary, lambda: ReportService.aging("sale"), ReportService.stock_valuation):
        with pytest.raises(PermissionDenied):
            call()
